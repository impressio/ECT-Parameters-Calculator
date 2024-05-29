"""Skapar transkript för talsyntes (ElevenLabs) ur talaranteckningarna i anteckningar.json.

Skriver mappen tal/: en textfil per bild med anteckning, och manus.txt med alla anteckningar i ordning,
åtskilda av pauser vars längd följer presentationens faser. Bildernas rubriker hämtas ur den byggda presentationen. Kör: python3 presentation/transkript.py"""
import json
import re
from pathlib import Path
from pptx import Presentation

HERE = Path(__file__).resolve().parent
OUT = HERE / 'tal'
# Pauser mellan bilderna i manus.txt, anpassade efter presentationens faser. Eleven v3 och v4 stöder inte SSML-taggen
# <break>; pausens längd styrs i stället av ljudtaggar och ellips (...), som saktar ner slutet av föregående mening.
SECTION = 'Startbild enfärgad'   # layouten för avsnittsbilderna (Stimulus och impedans, Dosering, Vid bordet, Säkerhet, Frågor)
SUMMARY = 'Rubrik, innehåll och ljusbrun bakgrund'   # layouten för sammanfattningen Att ta med sig
PAUSES = {
    'kort': '[pause]',                                    # mellan bilder inom samma avsnitt
    'mellan': '[long pause]',                             # efter titel, agenda och avsnittsbild samt före och efter sammanfattningen
    'lång': '[long pause]\n\n...\n\n[long pause]',            # före ett nytt avsnitt eller frågestunden
}
WITHIN = {'stycke': '[pause]', 'rad': '[short pause]'}   # inom en bild: mellan avsnitt (tom rad) och efter radbyte i anteckningen
# Tempo: Eleven v3 och v4 har inga fasta tempotaggar men följer beskrivande ljudtaggar direkt före texten de gäller
# (ElevenLabs, Prompting Eleven v4). Varje bild inleds med en tempotagg: lugnt tempo, och långsammare på bilder med
# mätvärden. Grundtempot ställs med inställningen Speed (0,7-1,2; under 1,0 talar rösten långsammare).
PACE = {'lugn': '[calm, measured pace]', 'långsam': '[slow, deliberate pace]'}
# Bilder med mätvärden (tal följt av en enhet) läses långsammare
UNITS = re.compile(r'\d[\d ,]*\s*(milli\w*|mikro\w*|sekund\w*|hertz|ohm|joule|procent|grad\w*|kilo\w*|gånger)', re.I)
OPENING = f"{PACE['lugn']} Elektrokonvulsiv terapi (ECT)...\n\n[long pause]"   # inledning före första bilden i manus.txt och i första bildens fil

notes = json.load(open(HERE / 'anteckningar.json', encoding='utf-8'))
slides = Presentation(HERE / 'ECT vid bordet.pptx').slides
assert len(notes) == len(slides), 'anteckningar.json och presentationen har olika antal bilder'
OUT.mkdir(exist_ok=True)
for old in OUT.glob('*.txt'):
    old.unlink()
parts = []  # (text, layout) för bilder med anteckning
def speech(note):
    """Tempotagg först; avsnitt (åtskilda av tom rad) får [pause], radbyten inom ett avsnitt [short pause]."""
    paragraphs = [p.strip() for p in note.strip().split('\n\n') if p.strip()]
    if not paragraphs:
        return ''
    pace = PACE['långsam'] if UNITS.search(note) else PACE['lugn']
    return pace + ' ' + f"\n\n{WITHIN['stycke']}\n\n".join(f"\n\n{WITHIN['rad']}\n\n".join(l.strip() for l in p.split('\n') if l.strip()) for p in paragraphs)


for i, (slide, note) in enumerate(zip(slides, notes), 1):
    note = speech(note)
    if not note:
        continue
    title = slide.shapes.title.text if slide.shapes.title is not None else f'Bild {i}'
    name = re.sub(r'\s+', ' ', re.sub(r'[^\w ,åäöÅÄÖ-]', '', title.replace('%', 'procent'))).strip()
    text = OPENING + '\n\n' + note if not parts else note
    (OUT / f'{i:02d} {name}.txt').write_text(text + '\n', encoding='utf-8')
    parts.append((note, slide.slide_layout.name))
manus = [OPENING]
for k, (note, layout) in enumerate(parts):
    if k == len(parts) - 1:
        manus.append(note)
        break
    next_layout = parts[k + 1][1]
    if next_layout == SECTION:
        pause = PAUSES['lång']
    elif layout in (SECTION, SUMMARY) or next_layout == SUMMARY or k < 2:
        pause = PAUSES['mellan']
    else:
        pause = PAUSES['kort']
    ending = note[:-1] + '...' if note.endswith('.') else note + ' ...'
    manus.append(ending + '\n\n' + pause)
(OUT / 'manus.txt').write_text('\n\n'.join(manus) + '\n', encoding='utf-8')
print(f'{len(parts)} bilder med anteckningar, {sum(len(p) for p, _ in parts)} tecken; skrivet till {OUT}')
