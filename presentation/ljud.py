"""Synkroniserar inspelat tal (ElevenLabs) med presentationen 'ECT vid bordet'.

Två steg:

  python3 presentation/ljud.py tidslinje TAL.wav --modell ggml-large-v3-turbo-q5_0.bin
      Transkriberar talet med whisper.cpp (whisper-cli) och letar upp var varje bilds anteckning börjar.
      Skärpunkterna läggs mitt i närmaste paus och sparas i tal/tidslinje.json.

  python3 presentation/ljud.py bygg TAL.wav
      Delar talet i ett klipp per bild (WAV, ljud/), bygger 'ECT vid bordet med ljud.pptx' där varje bild
      spelar sitt klipp automatiskt och går vidare när klippet är slut, och filmen 'ECT vid bordet.mp4'
      (bilderna i 1920 x 1080 med talet). Kräver ffmpeg, LibreOffice (soffice) och pdftoppm.

PowerPoint och LibreOffice Impress startar ljud på olika sätt, så varje bild får klippet två gånger (samma fil):
  - som dolt ljudobjekt som startar automatiskt (bildens tidslinje). Det fungerar i PowerPoint, även på första bilden
    och när man backar;
  - som övergångsljud (p:sndAc), som LibreOffice spelar.
Båda ligger i mc:AlternateContent-block vars val kräver PowerPoint 2016-namnrymden (p159), som PowerPoint förstår och
LibreOffice inte: ljudobjektet finns bara i valet och övergångsljudet bara i reservvarianten, så varje program spelar
klippet en gång.
    PowerPoint spelar övergångsljud vid fel tillfällen (inte på första bilden, och nästa bilds ljud när man backar).
Övergångsljud måste vara WAV; 22,05 kHz mono räcker för tal och håller nere storleken.

Inspelningen och klippen hålls utanför Git (se .gitignore)."""
import argparse
import difflib
import json
import math
import re
import subprocess
import sys
import tempfile
import zlib
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.media import Video
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.parts.image import Image, ImagePart
from pptx.parts.media import MediaPart

HERE = Path(__file__).resolve().parent
PPTX = HERE / 'ECT vid bordet.pptx'
NOTES = HERE / 'anteckningar.json'
TIMELINE = HERE / 'tal' / 'tidslinje.json'
CLIPS = HERE / 'ljud'
OUT_PPTX = HERE / 'ECT vid bordet med ljud.pptx'
OUT_MP4 = HERE / 'ECT vid bordet.mp4'
TAIL = 3.0           # sekunder som sista bilden (logotypen) visas efter talet i filmen
AUDIO_RATE = '192k'  # AAC-bithastighet för filmens ljud
CLIP_RATE = '22050'  # samplingsfrekvens för klippen i presentationen (WAV, 16 bitar, mono)
PML = 'http://schemas.openxmlformats.org/presentationml/2006/main'
MC = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
P14 = 'http://schemas.microsoft.com/office/powerpoint/2010/main'
P159 = 'http://schemas.microsoft.com/office/powerpoint/2015/09/main'
RT_MEDIA = 'http://schemas.microsoft.com/office/2007/relationships/media'


def run(cmd, **kw):
    return subprocess.run(cmd, check=True, **kw)


def duration(path):
    out = run(['ffprobe', '-v', 'error', '-show_entries', 'format=duration', '-of', 'csv=p=0', str(path)],
              capture_output=True, text=True).stdout
    return float(out)


def norm(text):
    return re.sub(r'[^a-zåäö0-9 ]', ' ', text.lower()).split()


def timeline(wav, model):
    """Transkriberar talet och sparar var varje bild börjar (sekunder)."""
    notes = json.load(open(NOTES, encoding='utf-8'))
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        run(['ffmpeg', '-v', 'error', '-y', '-i', str(wav), '-ar', '16000', '-ac', '1', '-c:a', 'pcm_s16le', str(tmp / 'tal.wav')])
        run(['whisper-cli', '-m', str(model), '-f', str(tmp / 'tal.wav'), '-l', 'sv', '-oj', '-of', str(tmp / 'tal')],
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        segments = json.load(open(tmp / 'tal.json'))['transcription']
        silence = subprocess.run(['ffmpeg', '-hide_banner', '-i', str(wav), '-af', 'silencedetect=noise=-40dB:d=0.7', '-f', 'null', '-'],
                                 capture_output=True, text=True).stderr
    pauses = [(float(e), float(d)) for e, d in re.findall(r'silence_end: ([\d.]+) \| silence_duration: ([\d.]+)', silence)]
    # Ordtidslinje: segmentets tid fördelas jämnt över dess ord.
    words = []
    for seg in segments:
        a, b = seg['offsets']['from'] / 1000, seg['offsets']['to'] / 1000
        w = norm(seg['text'])
        words += [(a + (b - a) * k / max(1, len(w)), x) for k, x in enumerate(w)]
    tokens = [x for _, x in words]
    total = duration(wav)
    cuts, pos = [], 0
    for i, note in enumerate(notes):
        if not note.strip():
            continue
        if not cuts:
            cuts.append(0.0)
            continue
        head = norm(note)[:8]
        best = (0.0, pos)
        for j in range(pos, min(len(tokens), pos + 2500)):
            r = difflib.SequenceMatcher(None, head, tokens[j:j + len(head)]).ratio()
            if r > best[0]:
                best = (r, j)
            if r == 1:
                break
        if best[0] < 0.5:
            sys.exit(f'Bild {i + 1}: hittar inte anteckningens början i talet ({best[0]:.2f}).')
        start = words[best[1]][0]
        near = [(abs(e - start), e, d) for e, d in pauses if -4 <= start - e <= 2.5]
        if near:
            _, end, length = min(near)
            cut = end - min(length / 2, 0.6)   # mitt i pausen, högst 0,6 s före talet
        else:
            cut = start - 0.3
        cuts.append(round(cut, 2))
        pos = best[1] + 1
    slides_with_notes = [i + 1 for i, n in enumerate(notes) if n.strip()]
    data = {'ljudfil': Path(wav).name, 'längd': round(total, 2), 'bilder': slides_with_notes, 'start': cuts}
    TIMELINE.write_text(json.dumps(data, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
    print(f'{len(cuts)} bilder i tidslinjen, {total:.1f} s; sparat i {TIMELINE}')


def split(wav, data):
    """Ett WAV-klipp per bild med tal."""
    CLIPS.mkdir(exist_ok=True)
    for old in list(CLIPS.glob('*.wav')) + list(CLIPS.glob('*.m4a')):
        old.unlink()
    ends = data['start'][1:] + [data['längd']]
    clips = {}
    for slide, a, b in zip(data['bilder'], data['start'], ends):
        path = CLIPS / f'{slide:02d}.wav'
        run(['ffmpeg', '-v', 'error', '-y', '-ss', f'{a:.3f}', '-to', f'{b:.3f}', '-i', str(wav),
             '-ar', CLIP_RATE, '-ac', '1', '-c:a', 'pcm_s16le', str(path)])
        clips[slide] = path
    return clips


def png_1x1():
    """Genomskinlig bild 1 x 1 bildpunkt som ikon för ljudobjektet."""
    def chunk(kind, data):
        return len(data).to_bytes(4, 'big') + kind + data + zlib.crc32(kind + data).to_bytes(4, 'big')
    return (b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', (1).to_bytes(4, 'big') * 2 + bytes([8, 6, 0, 0, 0]))
            + chunk(b'IDAT', zlib.compress(b'\x00\x00\x00\x00\x00')) + chunk(b'IEND', b''))


def add_audio(slide, media_rid, link_rid, image_rid, name, slide_width):
    """Dolt ljudobjekt utanför bilden som startar automatiskt när bilden visas (PowerPoint).

    Objektet ligger i ett mc:AlternateContent-val som kräver p159 och saknar reservvariant, så LibreOffice hoppar
    över det (annars spelar LibreOffice både objektet och övergångsljudet). Tidslinjens kommando pekar då inte på
    något objekt i LibreOffice."""
    tree = slide.shapes._spTree
    spid = max(int(e.get('id')) for e in tree.iter(qn('p:cNvPr'))) + 1
    pic = etree.fromstring(f'''<p:pic xmlns:p="{PML}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<p:nvPicPr><p:cNvPr id="{spid}" name="{name}"><a:hlinkClick r:id="" action="ppaction://media"/></p:cNvPr>
<p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr>
<p:nvPr><a:audioFile r:link="{link_rid}"/><p:extLst><p:ext uri="{{DAA4B4D4-6D71-4841-9C94-3DE7FCFB9230}}"><p14:media xmlns:p14="{P14}" r:embed="{media_rid}"/></p:ext></p:extLst></p:nvPr></p:nvPicPr>
<p:blipFill><a:blip r:embed="{image_rid}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>
<p:spPr><a:xfrm><a:off x="{slide_width + 100000}" y="0"/><a:ext cx="100000" cy="100000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>
</p:pic>''')
    alternate = etree.SubElement(tree, f'{{{MC}}}AlternateContent', nsmap={'mc': MC})
    choice = etree.SubElement(alternate, f'{{{MC}}}Choice', nsmap={'p159': P159})
    choice.set('Requires', 'p159')
    choice.append(pic)
    return spid


def timing(spid, ms):
    """Tidslinje: ljudobjektet startar direkt när bilden visas ("Starta automatiskt") och döljs i bildspelet."""
    return etree.fromstring(f'''<p:timing xmlns:p="{PML}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="{ms}" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
<p:audio><p:cMediaNode vol="80000" showWhenStopped="0"><p:cTn id="7" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:endCondLst><p:cond evt="onStopAudio" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:endCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:audio>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>''')


def set_advance(slide, ms, sound_rid, sound_name, spid):
    """Bilden går vidare automatiskt efter ms millisekunder; ljudobjektet spelas i PowerPoint, övergångsljudet i LibreOffice.

    Mallens egna övergångar ersätts. Ordningen i p:sld är cSld, clrMapOvr, transition (här i mc:AlternateContent), timing."""
    sld = slide._element
    for old in sld.findall(qn('p:timing')) + sld.findall(qn('p:transition')) + sld.findall(f'{{{MC}}}AlternateContent'):
        sld.remove(old)
    anchor = sld.find(qn('p:clrMapOvr'))
    anchor = anchor if anchor is not None else sld.find(qn('p:cSld'))
    alternate = etree.Element(f'{{{MC}}}AlternateContent', nsmap={'mc': MC})
    choice = etree.SubElement(alternate, f'{{{MC}}}Choice', nsmap={'p159': P159})
    choice.set('Requires', 'p159')
    fallback = etree.SubElement(alternate, f'{{{MC}}}Fallback')
    for parent, sound in ((choice, False), (fallback, True)):
        transition = etree.SubElement(parent, qn('p:transition'))
        transition.set('advTm', str(ms))
        # Utan angiven effekt ignorerar LibreOffice advTm; p:cut är en omedelbar växling utan synlig effekt.
        etree.SubElement(transition, qn('p:cut'))
        if sound:
            snd = etree.SubElement(etree.SubElement(etree.SubElement(transition, qn('p:sndAc')), qn('p:stSnd')), qn('p:snd'))
            snd.set(qn('r:embed'), sound_rid)
            snd.set('name', sound_name)
    anchor.addnext(alternate)
    alternate.addnext(timing(spid, ms))


def build_pptx(clips, data, tmp):
    prs = Presentation(PPTX)
    lengths = dict(zip(data['bilder'], [b - a for a, b in zip(data['start'], data['start'][1:] + [data['längd']])]))
    icon = ImagePart.new(prs.part.package, Image.from_blob(png_1x1(), 'ljud.png'))
    for number, slide in enumerate(prs.slides, 1):
        if number in clips:
            media = MediaPart.new(slide.part.package, Video.from_path_or_file_like(str(clips[number]), 'audio/wav'))
            link_rid = slide.part.relate_to(media, RT.AUDIO)
            media_rid = slide.part.relate_to(media, RT_MEDIA)
            image_rid = slide.part.relate_to(icon, RT.IMAGE)
            # Hela sekunder uppåt: LibreOffice avrundar växlingstiden till sekunder, och bilden får inte växla före talets slut.
            ms = math.ceil(lengths[number]) * 1000
            spid = add_audio(slide, media_rid, link_rid, image_rid, clips[number].name, prs.slide_width)
            set_advance(slide, ms, link_rid, clips[number].name, spid)
    prs.save(OUT_PPTX)
    print('sparad', OUT_PPTX)


def build_mp4(wav, data, tmp):
    run(['soffice', f'-env:UserInstallation=file://{tmp}/lo', '--headless', '--convert-to', 'pdf', '--outdir', str(tmp), str(PPTX)],
        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    run(['pdftoppm', '-png', '-scale-to-x', '1920', '-scale-to-y', '1080', str(tmp / (PPTX.stem + '.pdf')), str(tmp / 'bild')])
    images = sorted(tmp.glob('bild-*.png'))
    starts = dict(zip(data['bilder'], data['start']))
    ends = dict(zip(data['bilder'], data['start'][1:] + [data['längd']]))
    lines = []
    for number, image in enumerate(images, 1):
        seconds = ends[number] - starts[number] if number in starts else TAIL
        lines += [f"file '{image}'", f'duration {seconds:.3f}']
    lines.append(f"file '{images[-1]}'")   # concat-demuxern kräver att sista bilden upprepas
    (tmp / 'bilder.txt').write_text('\n'.join(lines) + '\n')
    total = data['längd'] + sum(TAIL for n in range(1, len(images) + 1) if n not in starts)
    run(['ffmpeg', '-v', 'error', '-y', '-f', 'concat', '-safe', '0', '-i', str(tmp / 'bilder.txt'), '-i', str(wav),
         '-af', 'apad', '-t', f'{total:.3f}', '-vf', 'fps=25,format=yuv420p', '-c:v', 'libx264', '-tune', 'stillimage',
         '-crf', '20', '-c:a', 'aac', '-b:a', AUDIO_RATE, '-movflags', '+faststart', str(OUT_MP4)])
    print('sparad', OUT_MP4, f'({duration(OUT_MP4):.1f} s)')


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest='steg', required=True)
    t = sub.add_parser('tidslinje')
    t.add_argument('tal')
    t.add_argument('--modell', required=True)
    b = sub.add_parser('bygg')
    b.add_argument('tal')
    args = parser.parse_args()
    if args.steg == 'tidslinje':
        timeline(Path(args.tal), Path(args.modell))
        return
    data = json.load(open(TIMELINE, encoding='utf-8'))
    if abs(duration(args.tal) - data['längd']) > 0.5:
        sys.exit('Ljudfilen har inte samma längd som tidslinjen; kör steget tidslinje igen.')
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        clips = split(Path(args.tal), data)
        build_pptx(clips, data, tmp)
        build_mp4(Path(args.tal), data, tmp)


if __name__ == '__main__':
    main()
