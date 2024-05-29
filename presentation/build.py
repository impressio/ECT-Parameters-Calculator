"""Bygger presentationen 'ECT vid bordet' på Sahlgrenskas mall.

Indata ligger i samma mapp: mall.pptx, bilder/ och anteckningar.json. Kör: python3 presentation/build.py [utfil]"""
import json
import sys
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from typsnitt import FONT, embed
from metadata import clean

HERE = Path(__file__).resolve().parent
TPL = HERE / 'mall.pptx'                 # Sahlgrenskas ursprungliga mall (Instruktionssida.pptx)
IMG = HERE / 'bilder'                    # bilder som bäddas in
NOTES_FILE = HERE / 'anteckningar.json'  # talaranteckningar, en text per bild i ordning
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / 'ECT vid bordet.pptx'
prs = Presentation(TPL)

# Ta bort mallens instruktions- och exempelbilder
lst = prs.slides._sldIdLst
for sid in list(lst):
    prs.part.drop_rel(sid.get(qn('r:id')))
    lst.remove(sid)

BLUE = MSO_THEME_COLOR.ACCENT_1          # 005B89, huvudfärg (vit text)
BROWN = MSO_THEME_COLOR.BACKGROUND_2     # E7E1DF, ljusbrun (svart text)
LIGHTBLUE = RGBColor(0xB3, 0xEB, 0xF2)   # ljusblå enligt mallens layouter (svart text)
BLACK = MSO_THEME_COLOR.TEXT_1
WHITE = MSO_THEME_COLOR.BACKGROUND_1

L_START, L_SECTION, L_CONTENT, L_BROWN, L_LIGHTBLUE, L_MAX = 0, 2, 4, 12, 13, 14


def add(layout, title):
    s = prs.slides.add_slide(prs.slide_layouts[layout])
    if layout == L_MAX:  # innehållet ritas som egna former; tom platshållare tas bort
        empty = ph(s, 14)
        empty._element.getparent().remove(empty._element)
    if title is not None:
        s.shapes.title.text = title
    return s


def ph(slide, idx):
    return next(p for p in slide.placeholders if p.placeholder_format.idx == idx)


def color(font_or_fill, c):
    target = font_or_fill.color if hasattr(font_or_fill, 'color') and not hasattr(font_or_fill, 'solid') else font_or_fill.fore_color
    if isinstance(c, RGBColor):
        target.rgb = c
    else:
        target.theme_color = c


def fill(shape, c):
    if c is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        color(shape.fill, c)
    shape.line.fill.background()
    shape.shadow.inherit = False


def write(tf, paras, size=16, fg=BLACK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=6, bullets=False):
    """paras: list of str or (str, dict(bold, size)) tuples."""
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    first = True
    for item in paras:
        text, opt = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        if not opt.get('nospace'):
            p.space_before = Pt(space)
        if bullets and not opt.get('bold'):
            pPr = p._p.get_or_add_pPr()
            pPr.set('marL', '269875')
            pPr.set('indent', '-269875')
            bf = etree.SubElement(pPr, qn('a:buFont'))
            bf.set('typeface', FONT)
            etree.SubElement(pPr, qn('a:buChar')).set('char', '•')
        r = p.add_run()
        r.text = text
        r.font.size = Pt(opt.get('size', size))
        r.font.bold = opt.get('bold', False)
        color(r.font, opt.get('fg', fg))
    return tf


def box(slide, x, y, w, h, bg, paras, shape=MSO_SHAPE.RECTANGLE, inset=0.18, **kw):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(sh, bg)
    tf = sh.text_frame
    for side in ('left', 'right', 'top', 'bottom'):
        setattr(tf, f'margin_{side}', Inches(inset))
    if 'fg' not in kw:
        kw['fg'] = WHITE if bg == BLUE else BLACK
    write(tf, paras, **kw)
    return sh


def circle(slide, x, y, d, bg, paras, **kw):
    kw.setdefault('align', PP_ALIGN.CENTER)
    kw.setdefault('anchor', MSO_ANCHOR.MIDDLE)
    return box(slide, x, y, d, d, bg, paras, shape=MSO_SHAPE.OVAL, inset=0.05, **kw)


def text(slide, x, y, w, h, paras, **kw):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = Inches(0)
    write(tf, paras, **kw)
    return tb


def source(slide, s):
    text(slide, 0.39, 6.82, 12.55, 0.35, [(s, {'nospace': True})], size=11)


def body(slide, idx, paras, size=20, h=None, w=None):
    p = ph(slide, idx)
    p.left, p.top, p.width, p.height = p.left, p.top, p.width, p.height
    if h is not None:
        p.height = Inches(h)
    if w is not None:
        p.width = Inches(w)
    tf = p.text_frame
    first = True
    for item in paras:
        t, opt = (item, {}) if isinstance(item, str) else (item[0], item[1] if len(item) > 1 else {})
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if isinstance(t, tuple):  # (fet inledning, resten)
            r = para.add_run(); r.text = t[0]; r.font.bold = True; r.font.size = Pt(size)
            r = para.add_run(); r.text = t[1]; r.font.size = Pt(size)
        else:
            r = para.add_run(); r.text = t; r.font.size = Pt(opt.get('size', size))
            r.font.bold = opt.get('bold', False)
            if opt.get('nobullet'):
                pPr = para._p.get_or_add_pPr()
                pPr.set('marL', '0'); pPr.set('indent', '0')
                etree.SubElement(pPr, qn('a:buNone'))


def section(title, sub):
    s = add(L_SECTION, title)
    ph(s, 1).text_frame.text = sub


def chart_text(chart, size=14):
    chart.font.size = Pt(size)
    chart.font.name = FONT
    color(chart.font, BLACK)


def no_line(el):
    spPr = el.find(qn('c:spPr'))
    if spPr is None:
        spPr = etree.SubElement(el, qn('c:spPr'))
    etree.SubElement(spPr, qn('a:ln')).append(etree.Element(qn('a:noFill')))


# 1 Titel
s = add(L_START, 'ECT vid bordet')
ph(s, 1).text_frame.text = 'Stimulus, impedans, anfall och säkerhet'

# 2 Agenda
s = add(L_LIGHTBLUE, '30 minuter och frågestund')
body(s, 17, [
    (('Inledning, 4,5 min: ', 'historia, indikationer och verkningsmekanismer'),),
    (('Stimulus & impedans, 8 min: ', 'apparat, parameterinställning, elektroder och hud'),),
    (('Dosering, 3 min: ', 'kramptröskel, multipel och strömstyrka'),),
    (('Vid bordet, 5,5 min: ', 'steg för steg, motorik och EEG, avvikande anfall'),),
    (('Säkerhet, 2 min: ', 'läkemedel och biverkningar'),),
    ('Därefter film och frågestund', {'bold': True, 'nobullet': True}),
], size=20, h=4.3)

# Historia
s = add(L_MAX, 'ECT:s historia i korthet')
line = s.shapes.add_connector(1, Inches(0.9), Inches(2.3), Inches(12.4), Inches(2.3))
line.line.color.theme_color = BLUE
line.line.width = Pt(2)
history = [('1934', 'Anfall med kamfer och kardiazol', 'Meduna, Budapest'),
           ('1938', 'Första elektriskt utlösta anfallet', 'Cerletti och Bini, Rom'),
           ('1952', 'Suxameton ger muskelrelaxation', 'Holmberg och Thesleff, Sverige'),
           ('1958', 'Unilateral ECT', 'Lancaster m.fl.'),
           ('1987', 'Dosering efter kramptröskeln; kortpuls ersätter sinusvåg', 'Sackeim m.fl.')]
for i, (year, head, who) in enumerate(history):
    x = 0.39 + i * 2.53
    circle(s, x + 0.6, 1.7, 1.2, BLUE, [(year, {'bold': True, 'size': 20, 'nospace': True})])
    text(s, x, 3.05, 2.4, 1.9, [(head, {'bold': True, 'nospace': True}), who], size=16, align=PP_ALIGN.CENTER)
box(s, 0.39, 5.25, 12.55, 1.3, LIGHTBLUE, [('Sverige 2013: knappt 4000 patienter, 41 per 100 000 invånare; '
                                           '86 % av serierna gavs unilateralt', {'nospace': True})],
    size=18, anchor=MSO_ANCHOR.MIDDLE, inset=0.25)
source(s, 'Källor: Endler 1988, Convuls Ther; Baran m.fl. 2012, J Affect Disord; Holmberg och Thesleff 1952, Am J Psychiatry; '
          'Lancaster m.fl. 1958, J Ment Sci; Sackeim m.fl. 1987; Nordanskog m.fl. 2015, J ECT')

# Indikationer och kontraindikationer
s = add(L_MAX, 'Indikationer och kontraindikationer')
box(s, 0.39, 1.75, 5.3, 4.8, BLUE, [('Indikationer', {'bold': True, 'size': 22, 'nospace': True}),
                                    'Svår depression, särskilt psykotisk eller melankolisk',
                                    'Katatoni, malignt neuroleptikasyndrom',
                                    'Akut suicidrisk, vägrar äta och dricka',
                                    'Svår eller behandlingsresistent mani',
                                    'Behandlingsresistent depression',
                                    'Schizofreni med otillräcklig läkemedelseffekt'],
    size=16, space=8, bullets=True, inset=0.3)
box(s, 5.84, 1.75, 7.1, 1.0, LIGHTBLUE, [('Absoluta kontraindikationer: inga', {'bold': True, 'size': 20, 'nospace': True})],
    anchor=MSO_ANCHOR.MIDDLE, inset=0.3)
box(s, 5.84, 2.9, 7.1, 3.65, BROWN, [('Relativa kontraindikationer', {'bold': True, 'size': 22, 'nospace': True}),
                                    'Hjärtinfarkt senaste 3 mån, instabil angina',
                                    'Dekompenserad hjärtsvikt',
                                    'Svår aortastenos',
                                    'Aneurysm eller arteriovenös missbildning',
                                    'Förhöjt intrakraniellt tryck',
                                    'Stroke eller hjärnblödning senaste 1-3 mån',
                                    'Feokromocytom, näthinneavlossning'],
    size=16, space=6, bullets=True, inset=0.3)
source(s, 'Källor: APA 2025, kap. 2.2-2.4 och 9.2; The ECT Handbook, kap. 3-5, 10-11 och 27; Mankad m.fl. 2010, kap. 2-3; SPF:s riktlinjer 2014')

# Verkningsmekanismer
s = add(L_MAX, 'Verkningsmekanismer och neurobiologiska korrelat')
mech = [('Anfallet krävs', 'Generaliserat anfall; subkonvulsiv stimulering ger svagare effekt', BROWN),
        ('Specifika nätverk', 'RUL aktiverar asymmetriskt; hjärnstammen alltid symmetriskt', BROWN),
        ('Signalsubstanser', 'Monoaminer och GABA/glutamat; ökad hämmande kapacitet', BROWN),
        ('HPA-axeln', 'Stresshormonaxeln är rubbad vid svår depression; ECT normaliserar den vid melankoli', BLUE),
        ('Hippocampus', 'Ökad cellnybildning i djurstudier; betydelsen hos människa oklar', BROWN),
        ('Funktionella nätverk', 'ECT påverkar nätverk som är i obalans vid depression', BROWN)]
for i, (head, t, bg) in enumerate(mech):
    x, y = 0.39 + (i % 3) * 4.23, 1.75 + (i // 3) * 2.45
    box(s, x, y, 4.03, 2.25, bg, [(head, {'bold': True, 'size': 20, 'nospace': True}), t], size=16, space=10, inset=0.25)
source(s, 'Källor: SPF:s riktlinjer 2014, avsnitt 2.4, s. 21-23; The ECT Handbook, kap. 2')

# 3
section('Stimulus och impedans', 'Del 1 av 4  ·  8 minuter')

# Apparaten
s = add(L_MAX, 'Apparaten: MECTA spECTrum 5000Q')
pic = s.shapes.add_picture(str(IMG / 'mecta-5000q-foto.jpg'), Inches(0.39), Inches(1.75), width=Inches(7.6))
pic.name = 'MECTA spECTrum 5000Q, foto'
pic._element.nvPicPr.cNvPr.set('descr', 'Foto av MECTA spECTrum 5000Q med skärm, parameterinställningar, EEG-ingångar, skrivare, kablar och bruksanvisning')
for i, (head, t) in enumerate([
        ('Stimulus', 'Konstant ström 500-900 mA, bifasiska kortpulser; högst 1152 mC'),
        ('Impedans', 'Mäts före stimuleringen; spärr under 100 och över 5000 Ω'),
        ('EEG och skrivare', 'Två EEG-kanaler; pappersremsa efter behandlingen')]):
    box(s, 8.2, 1.75 + i * 1.62, 4.74, 1.47, BROWN, [(head, {'bold': True, 'size': 18, 'nospace': True}), t],
        size=16, space=4, inset=0.2, anchor=MSO_ANCHOR.MIDDLE)
source(s, 'Källor: MECTA spECTrum 5000Q, bruksanvisning och datablad (200 J); Mankad m.fl. 2010, kap. 4')

# 4 Parameterinställning
s = add(L_MAX, 'Parameterinställningen bestämmer laddningen')
pic = s.shapes.add_picture(str(IMG / 'mecta-5000q-frontpanel.png'), Inches(0.39), Inches(1.75), width=Inches(7.7))
pic.name = 'MECTA spECTrum 5000Q'
pic._element.nvPicPr.cNvPr.set('descr', 'Frontpanel på MECTA spECTrum 5000Q med parameterinställningarna Pulse width, Frequency, Duration och Current')
for i, (name, rng) in enumerate([('Pulsbredd', '0,3-1,0 ms'), ('Frekvens', '20-120 Hz'),
                                 ('Duration', '0,5-8,0 s'), ('Strömstyrka', '500-900 mA')]):
    box(s, 0.39 + i * 1.97, 4.95, 1.82, 1.15, BROWN,
        [(name, {'bold': True, 'nospace': True}), rng], size=14, inset=0.12)
box(s, 8.45, 1.75, 4.49, 4.35, BLUE, [
    ('Laddning (mC)', {'bold': True, 'nospace': True}),
    ('Q = PW × f × 2 × D × I / 1000', {'size': 20, 'bold': True}),
    ('1,0 ms × 60 Hz × 2 × 2,0 s × 800 mA / 1000 = 192 mC', {}),
    ('Bifasiska pulser: en positiv och en negativ per period', {'size': 14}),
], size=16, space=16, inset=0.3, anchor=MSO_ANCHOR.MIDDLE)
source(s, 'Källor: MECTA spECTrum 5000Q, bruksanvisning; Mankad m.fl. 2010, kap. 4')

# 5 1 %
s = add(L_MAX, 'Bara omkring 1 % av strömmen når hjärnan')
for i, (val, lab, bg, sz) in enumerate([
        ('~1 %', 'passerar skallbenet in i hjärnan', BLUE, 40),
        ('~99 %', 'leds genom hud och skalp och blir värme', BROWN, 40),
        ('< 0,001 °C', 'temperaturökning i hjärnan', LIGHTBLUE, 26)]):
    cx = 0.39 + i * 4.18
    circle(s, cx + 0.84, 1.7, 2.4, bg, [(val, {'bold': True, 'size': sz, 'nospace': True})])
    text(s, cx, 4.2, 4.0, 0.8, [(lab, {'nospace': True})], size=16, align=PP_ALIGN.CENTER)
box(s, 0.39, 5.15, 12.55, 1.35, BROWN, [
    ('Brännskada uppstår bara där kontakten är dålig: hög impedans i kontaktytan mellan elektrod och hud. '
     'Skallbenets tjocklek och huvudets anatomi avgör hur mycket som når hjärnan.', {'nospace': True})],
    size=16, anchor=MSO_ANCHOR.MIDDLE, inset=0.25)
source(s, 'Källor: Swartz 2009, kap. 1, s. 4-5; Fridgeirsson m.fl. 2021; Deng m.fl. 2022')

# 6 Impedans
s = add(L_MAX, 'Impedansen visar kontakten')
cards = [
    ('Statisk, före stimulering', ['Spärrfönster 100-5000 Ω', 'Målvärde omkring 1000 Ω'], BROWN),
    ('Dynamisk, under pulståget', ['Typiskt 150-350 Ω', 'Lägre än den statiska: huden leder bättre under pulserna'], BROWN),
    ('OVER: dålig kontakt', ['Hår, hudfett, för torr kompress, för löst tryck'], BLUE),
    ('UNDER: salt- eller gelbrygga', ['Fuktig yta mellan elektroderna: strömmen går i huden'], BLUE),
]
for i, (head, lines, bg) in enumerate(cards):
    x, y = 0.39 + (i % 2) * 6.35, 1.75 + (i // 2) * 2.5
    box(s, x, y, 6.2, 2.3, bg, [(head, {'bold': True, 'size': 20, 'nospace': True})] + lines,
        size=18, space=12, bullets=True, inset=0.25)
source(s, 'Källor: MECTA spECTrum 5000Q, bruksanvisning; Mankad m.fl. 2010, kap. 4; målvärdet: klinisk erfarenhet')

# Elektrodplacering
s = add(L_MAX, 'Elektrodplacering: stimulering och EEG')
for image, x, w, descr in [('elektrodplacering-stimulering.png', 1.14, 4.7, 'Huvuden i halvprofil med elektroderna T och V vid högersidig unilateral placering och T på båda sidor vid bitemporal placering'),
                           ('elektrodplacering-eeg.png', 8.55, 2.57, 'Huvud och bröst framifrån med två EEG-elektroder i pannan och referenselektroden under nyckelbenet')]:
    pic = s.shapes.add_picture(str(IMG / image), Inches(x), Inches(1.55), width=Inches(w))
    pic.name = image.rsplit('.', 1)[0]
    pic._element.nvPicPr.cNvPr.set('descr', descr)
box(s, 0.39, 4.7, 3.03, 1.95, BLUE, [('RUL d’Elia, förstahandsval', {'bold': True, 'size': 15, 'nospace': True}),
                                     'T: 2,5-3 cm ovanför mitten av linjen ögonvrå-tragus', 'V: 2,5-3 cm höger om hjässan; minst 10-12 cm mellan'],
    size=13, space=4, inset=0.15)
box(s, 3.56, 4.7, 3.03, 1.95, BROWN, [('Bitemporal (BT)', {'bold': True, 'size': 15, 'nospace': True}),
                                      'T på båda sidor, 2,5-3 cm ovanför mitten av linjen ögonvrå-tragus'],
    size=13, space=4, inset=0.15)
box(s, 6.74, 4.7, 6.2, 1.95, LIGHTBLUE, [('EEG, tre elektroder', {'bold': True, 'size': 15, 'nospace': True}),
                                         'Två i pannan, 2-3 cm ovanför ögonbrynen', 'Referens under nyckelbenet'],
    size=13, space=4, inset=0.15)
source(s, 'Källor: APA 2025, kap. 14.3; The ECT Handbook, kap. 22; d’Elia 1970; lokal rutin')

# 7 Hudförberedelse
s = add(L_MAX, 'Hudförberedelse: ren hud och saltkompress')
steps = [('Ren hud', 'Rent hår utan hårprodukter, osminkad'),
         ('NaCl-kompress', 'NaCl 250 ml + ½ msk salt, små rena kompresser; gel som alternativ'),
         ('Får inte rinna', 'Fuktig kompress, ingen vätska mellan elektroderna: risk för saltbrygga'),
         ('Ingen sprit', 'Används inte hos oss; måste vara helt torr: brandrisk nära syrgas')]
for i, (head, t) in enumerate(steps):
    x = 0.39 + i * 3.17
    circle(s, x, 1.85, 1.1, BLUE, [(str(i + 1), {'bold': True, 'size': 28, 'nospace': True})])
    box(s, x, 3.2, 3.0, 3.2, BROWN, [(head, {'bold': True, 'size': 22, 'nospace': True}), t],
        size=18, space=12, inset=0.22)
source(s, 'Källor: APA 2025, kap. 14.5; Swartz 2009, s. 5; lokal rutin')

# 8 Pulsbredd: effekt och minne
s = add(L_MAX, 'Pulsbredd: effekt mot minnesbiverkningar')
box(s, 0.39, 1.75, 6.2, 3.55, BLUE, [('Kortpuls, 0,5-1,0 ms', {'bold': True, 'size': 22, 'nospace': True}),
                                     'Standard hos oss och vanligast i Sverige',
                                     'Bäst underbyggd effekt',
                                     'Svenska registret: puls minst 0,5 ms gav oftare remission'],
    size=18, space=12, bullets=True, inset=0.3)
box(s, 6.74, 1.75, 6.2, 3.55, BROWN, [('Ultrakort puls, 0,25-0,3 ms', {'bold': True, 'size': 22, 'nospace': True}),
                                      'Mindre minnesbiverkningar vid RUL',
                                      'Något sämre effekt och i genomsnitt en behandling till',
                                      'Används inte hos oss'],
    size=18, space=12, bullets=True, inset=0.3)
box(s, 0.39, 5.5, 12.55, 1.05, LIGHTBLUE, [('Samma laddning (mC) är inte likvärdig vid olika pulsbredd: '
                                           'byte av pulsbredd kräver ny bedömning av dosen', {'nospace': True})],
    size=16, anchor=MSO_ANCHOR.MIDDLE, inset=0.25)
source(s, 'Källor: Sackeim m.fl. 2008, Brain Stimul; Tor m.fl. 2015, J Clin Psychiatry; Brus m.fl. 2017, Eur Psychiatry och J ECT; APA 2025, kap. 14')

# 9 Pulsbredd i Sverige
s = add(L_MAX, 'I Sverige är 0,5 ms vanligast')
cd = CategoryChartData()
cd.categories = ['0,25-0,47 ms', '0,50 ms', '0,51-1,00 ms']
cd.add_series('Andel (%)', (19.0, 68.4, 12.6))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.39), Inches(1.65),
                        Inches(8.1), Inches(5.05), cd)
gf.name = 'Pulsbredd i Sverige'
ch = gf.chart
chart_text(ch)
ch.has_legend = False
ch.has_title = False
plot = ch.plots[0]
plot.gap_width = 60
bars = plot.series[0]
bars.format.fill.solid(); bars.format.fill.fore_color.theme_color = BLUE
for i, lab in enumerate(['19,0 %', '68,4 %', '12,6 %']):
    dl = bars.points[i].data_label
    dl.text_frame.text = lab
    dl.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
    dl.text_frame.paragraphs[0].runs[0].font.bold = True
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
va = ch.value_axis
va.minimum_scale, va.maximum_scale, va.major_unit = 0, 80, 20
va.has_major_gridlines = True
va.major_gridlines.format.line.color.theme_color = BROWN
va.has_title = True
va.axis_title.text_frame.text = 'Andel patienter (%)'
r = va.axis_title.text_frame.paragraphs[0].runs[0]
r.font.size = Pt(14); r.font.bold = False; r.font.name = FONT
va.format.line.fill.background()
ch.category_axis.format.line.color.rgb = RGBColor(0xA6, 0xA6, 0xA6)
box(s, 8.8, 1.75, 4.14, 2.3, BLUE, [('68,4 %', {'bold': True, 'size': 32, 'nospace': True}),
                                     'fick 0,50 ms (1671 patienter, 2011-2014)'], size=16, inset=0.25)
box(s, 8.8, 4.25, 4.14, 2.3, BROWN, [('Ultrakort puls (under 0,5 ms) används inte hos oss',
                                      {'nospace': True})], size=16, inset=0.25)
source(s, 'Källa: Brus m.fl. 2017, Eur Psychiatry (tabell 1)')

# 10
section('Dosering', 'Del 2 av 4  ·  3 minuter')

# 11 Kramptröskeln
s = add(L_MAX, 'Kramptröskeln styr dosen')
box(s, 0.39, 1.75, 5.4, 4.75, BROWN, [('Höjer tröskeln', {'bold': True, 'size': 20, 'nospace': True}),
                                      'Hög ålder', 'Manligt kön', 'Bensodiazepiner och antiepileptika',
                                      'Djup anestesi (propofol)',
                                      'Bitemporal placering (jämfört med RUL)'],
    size=18, space=12, bullets=True, inset=0.25)
rows = [('Placering och puls', 'Multipel'), ('RUL, kortpuls', '4-6 × tröskeln'),
        ('Bitemporal, kortpuls', '1,5-2,5 × tröskeln')]
gf = s.shapes.add_table(3, 2, Inches(6.05), Inches(1.75), Inches(6.89), Inches(1.95))
gf.name = 'Multipel av tröskeln'
tbl = gf.table
tbl.columns[0].width, tbl.columns[1].width = Inches(3.9), Inches(2.99)
for r_i, row in enumerate(rows):
    tbl.rows[r_i].height = Inches(0.65)
    for c_i, val in enumerate(row):
        cell = tbl.cell(r_i, c_i)
        cell.fill.solid()
        if r_i == 0:
            cell.fill.fore_color.theme_color = BLUE
        elif r_i % 2:
            cell.fill.fore_color.theme_color = WHITE
        else:
            cell.fill.fore_color.theme_color = BROWN
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.18)
        cell.text_frame.text = val
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(16)
        run.font.bold = r_i == 0
        color(run.font, WHITE if r_i == 0 else BLACK)
box(s, 6.05, 3.95, 6.89, 2.55, LIGHTBLUE, [('Räkneexempel', {'bold': True, 'nospace': True}),
                                          'RUL: tröskel 50 mC × 5 = 250 mC',
                                          'Bitemporal: tröskel 75 mC × 2 = 150 mC',
                                          ('Högre tröskel betyder inte högre dos', {'bold': True})],
    size=16, space=12, bullets=True, anchor=MSO_ANCHOR.MIDDLE, inset=0.25)
source(s, 'Källor: Sackeim m.fl. 1987, Arch Gen Psychiatry; APA 2025, kap. 14; Mankad m.fl. 2010, tabell 5-1')

# 12 Strömstyrka
s = add(L_MAX, 'Strömstyrkan ställs in enligt ordination')
circle(s, 0.9, 1.75, 3.0, BLUE, [('500-900 mA', {'bold': True, 'size': 28, 'nospace': True})])
text(s, 0.39, 4.95, 4.0, 1.5, [('Enligt ordination; dosen anpassas i första hand med laddningen', {'nospace': True})],
     size=16, align=PP_ALIGN.CENTER)
for i, (head, t) in enumerate([
        ('Anatomin', 'Skallbenets tjocklek och huvudets storlek avgör hur mycket som når hjärnan, inte kroppsvikten'),
        ('Lägre ström', '700-800 mA gav bättre effekt än 600 mA; verbalt flöde försämrades vid 800 mA (Abbott 2021)'),
        ('I Sverige 2011-2014', '528-800 mA hos 55,5 %, 850-890 mA hos 12,7 % och 900-930 mA hos 31,7 % (Brus 2017)')]):
    box(s, 4.8, 1.75 + i * 1.62, 8.14, 1.47, BROWN, [(head, {'bold': True, 'size': 18, 'nospace': True}), t],
        size=16, space=4, inset=0.2, anchor=MSO_ANCHOR.MIDDLE)
source(s, 'Källor: MECTA spECTrum 5000Q, datablad (200 J); Abbott m.fl. 2021; Brus m.fl. 2017, Eur Psychiatry (tabell 1); Deng m.fl. 2022; SPF:s riktlinjer 2014')

# 13
section('Vid bordet', 'Del 3 av 4  ·  5,5 minuter')

# 14 Steg för steg
s = add(L_MAX, 'Behandlingen steg för steg')
flow = [('Kontroller', 'ID, fasta, mediciner, blåstömning, blodtryck'), ('Förberedelse', 'Uppkoppling, hud, NaCl-kompress'),
        ('Induktion', 'Preoxygenering, propofol'), ('Relaxation', 'Suxameton för muskelrelaxation'),
        ('Klartecken', 'Bitblock, anestesins klartecken'), ('Stimulering', 'Elektroder och impedans, håll knappen tills tonen tystnar')]
for i, (head, t) in enumerate(flow):
    x = 0.39 + i * 2.1
    circle(s, x, 1.75, 0.9, BLUE, [(str(i + 1), {'bold': True, 'size': 24, 'nospace': True})])
    box(s, x, 2.85, 1.97, 2.25, BROWN, [(head, {'bold': True, 'size': 16, 'nospace': True}), t],
        size=14, space=8, inset=0.14)
box(s, 0.39, 5.25, 12.55, 1.3, LIGHTBLUE, [('Blodtrycket mäts före behandlingen och omedelbart efter att EEG-aktiviteten har upphört',
                                           {'bold': True}),
                                          'Ordinarie antihypertensiva och PPI ges på avdelningen före ECT, annars enligt anestesiplanen'],
    size=16, space=6, bullets=True, anchor=MSO_ANCHOR.MIDDLE, inset=0.25)
source(s, 'Källor: APA 2025, kap. 14; The ECT Handbook, kap. 22; lokal rutin')

# Motorisk kramp
s = add(L_MAX, 'Motorisk kramp: faser och avläsning')
phases = [('Under pulståget', 'Käke och nacke drar ihop sig av strömmen; hör inte till anfallet'),
          ('Tonisk fas', 'Ihållande kontraktion direkt eller inom några sekunder, vid tröskelnära stimulering upp till 10 s'),
          ('Klonisk fas', 'Rytmiska ryckningar som blir glesare; under relaxation ofta bara i tår, fingrar eller ansikte'),
          ('Motoriskt slut', 'När ryckningarna har upphört helt: läs av sekunderna på apparaten')]
for i, (head, t) in enumerate(phases):
    x = 0.39 + i * 3.17
    circle(s, x, 1.75, 0.9, BLUE, [(str(i + 1), {'bold': True, 'size': 24, 'nospace': True})])
    box(s, x, 2.85, 3.0, 2.45, BROWN, [(head, {'bold': True, 'size': 18, 'nospace': True}), t], size=16, space=8, inset=0.2)
box(s, 0.39, 5.5, 12.55, 1.05, LIGHTBLUE, [('Motorisk duration vanligen över 20-30 s; EEG-anfallet varar vanligen 10-20 s längre. '
                                           'Sekundgränserna är inga absoluta effektkriterier.', {'nospace': True})],
    size=16, anchor=MSO_ANCHOR.MIDDLE, inset=0.25)
source(s, 'Källor: VGR:s regionala riktlinje ECT 2024; Mankad m.fl. 2010, kap. 7; The ECT Handbook, kap. 26, s. 234-236; APA 2025, kap. 15')

# EEG
s = add(L_MAX, 'EEG: ett anfall på pappersremsan')
rows = [(0.39, 1.7), (0.39, 3.23), (0.39, 4.76), (6.74, 1.7), (6.74, 3.23)]
for i, (x, y) in enumerate(rows):
    pic = s.shapes.add_picture(str(IMG / f'eeg-remsa-rad-{i + 1}.png'), Inches(x), Inches(y), width=Inches(6.2))
    pic.name = f'EEG-remsa, rad {i + 1}'
    pic._element.nvPicPr.cNvPr.set('descr', f'Klinisk tvåkanalig EEG-registrering under ECT, rad {i + 1} av 5'
                                   + (' (start till 8 s)', ' (11-19 s)', ' (24-32 s)', ' (37-45 s)', ' (46-53 s)')[i])
legend_box = box(s, 6.74, 4.76, 6.2, 1.4, BROWN, [])
legend_box.name = 'Teckenförklaring'
for col, items in enumerate([[('B', 'rekrytering'), ('C', 'polyspikar'), ('D', 'theta och delta'), ('E', 'spik-våg, 2-3 Hz')],
                             [('F', 'anfallet avtar'), ('G', 'suppression'), ('H', 'rörelseartefakt')]]):
    tb = s.shapes.add_textbox(Inches(6.95 + col * 3.05), Inches(4.86), Inches(2.95), Inches(1.25))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = Inches(0)
    tf.word_wrap = True
    for j, (letter, t) in enumerate(items):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if j == 0 else 2)
        r1 = p.add_run(); r1.text = letter + '  '; r1.font.bold = True; r1.font.size = Pt(16); color(r1.font, BLACK)
        r2 = p.add_run(); r2.text = t; r2.font.size = Pt(16); color(r2.font, BLACK)
source(s, 'Bild: Semple och Devlin, The ECT Handbook, 4 uppl. 2019, kap. 26, figur 26.7, s. 240. Källor: APA 2025, kap. 15; Mankad m.fl. 2010, kap. 7')

# 15 Två tider
s = add(L_MAX, 'Anfallets längd: motorisk kramp och EEG-tid')
for i, (head, lines) in enumerate([
        ('Motorisk duration', ['Från stimuleringen tills ryckningarna har upphört',
                               'Manschettmetoden används inte rutinmässigt']),
        ('EEG-tid', ['Från stimuleringen tills EEG planar ut',
                     'Verifiera slutet på EEG även när motoriken upphört'])]):
    box(s, 0.39 + i * 6.35, 1.65, 6.2, 1.75, BROWN, [(head, {'bold': True, 'size': 18, 'nospace': True})] + lines,
        size=16, space=8, bullets=True, inset=0.2)
pic = s.shapes.add_picture(str(IMG / 'pappersremsa.png'), Inches(0.39), Inches(3.65), width=Inches(10.6))
pic.name = 'Pappersremsa'
pic._element.nvPicPr.cNvPr.set('descr', 'Utskrift från MECTA spECTrum 5000Q med SELF TEST RESULTS, EEG-spår och TREATMENT RESULTS')
text(s, 11.2, 3.65, 1.74, 2.67, [('Utskrift direkt efter stimuleringen: SELF TEST RESULTS, EEG-spår och '
                                'TREATMENT RESULTS. Tiden läses av i sekunder.', {'nospace': True})],
     size=14)
source(s, 'Källor: The ECT Handbook, kap. 26, s. 234-236; MECTA-manualen (100 J), s. 57; lokal rutin')

# 16 Missat, kort, förlängt
s = add(L_MAX, 'Missat, kort eller förlängt anfall')
for i, (head, lines) in enumerate([
        ('Missat anfall', ['Inget anfall på EEG och ingen motorik',
                           'Avvakta minst 20 sekunder: anfallet kan starta sent',
                           'Kontrollera kontakt, brygga och levererad laddning',
                           'Ny stimulering är ett nytt beslut']),
        ('Kort anfall', ['Inte automatiskt ineffektivt: bedöm EEG',
                         'Restimulera inte bara för att en sekundgräns underskrids',
                         'Före ny stimulering: vänta 40-60 sekunder med ventilation (APA: minst 45)']),
        ('Förlängt anfall', ['Kramp längre än 1 minut (APA: omkring 2 minuter)',
                             'Informera anestesin; krampen bryts med diazepam (Stesolid) iv',
                             'Verifiera slutet på EEG'])]):
    box(s, 0.39 + i * 4.23, 1.75, 4.03, 4.8, BLUE if i == 2 else BROWN,
        [(head, {'bold': True, 'size': 22, 'nospace': True})] + lines,
        size=16, space=12, bullets=True, inset=0.25)
source(s, 'Källor: VGR:s regionala riktlinje ECT 2024; SPF:s riktlinjer 2014, avsnitt 7.6; APA 2025, kap. 21.1-21.3; IVA Mölndal, rutin ECT 2025')

# 17
section('Säkerhet', 'Del 4 av 4  ·  2 minuter')

# 18 Läkemedel
s = add(L_MAX, 'Läkemedel att stämma av')
meds = [('Bensodiazepiner, antiepileptika', 'Höjer tröskeln: korta eller missade anfall'),
        ('Litium', 'Delirium, förlängd blockad, sällsynt serotonergt syndrom'),
        ('Teofyllin', 'Förlängda anfall'),
        ('Kolinesterashämmare', 'Förlängd suxametoneffekt, bradykardi'),
        ('SSRI, SNRI', 'Sällsynt serotonergt syndrom'),
        ('Antihypertensiva, PPI', 'Ges på avdelningen före ECT')]
for i, (head, t) in enumerate(meds):
    x, y = 0.39 + (i % 3) * 4.23, 1.75 + (i // 3) * 2.45
    box(s, x, y, 4.03, 2.25, BLUE if i == 5 else BROWN,
        [(head, {'bold': True, 'size': 20, 'nospace': True}), t], size=18, space=10, inset=0.25)
source(s, 'Källor: The ECT Handbook, kap. 25; lokal rutin')

# 19 Biverkningar
s = add(L_MAX, 'Biverkningar och hur vi minskar dem')
effects = [('Minnet', 'RUL i första hand och två behandlingar i veckan'),
           ('Huvudvärk, muskelvärk, illamående', 'Vanligt; paracetamol, NSAID, antiemetika enligt ordination'),
           ('Postiktal delirium och agitation', 'Luftväg, syresättning och cirkulation först; lugn miljö'),
           ('Serotonergt syndrom', 'Sällsynt: klonus, hyperreflexi, tremor, svettning, feber')]
for i, (head, t) in enumerate(effects):
    x, y = 0.39 + (i % 2) * 6.35, 1.75 + (i // 2) * 2.45
    box(s, x, y, 6.2, 2.25, BROWN, [(head, {'bold': True, 'size': 22, 'nospace': True}), t],
        size=18, space=12, inset=0.3)
source(s, 'Källor: The ECT Handbook, kap. 13; Boyer och Shannon 2005; Dunkley m.fl. 2003')

# 20 Att ta med sig
s = add(L_BROWN, 'Att ta med sig')
body(s, 17, ['Skärmen ska visa ordinerad laddning i mC',
             'God kontakt: ren hud, NaCl-kompress som inte rinner, ingen brygga',
             'Dosen anpassas med laddningen; strömstyrka enligt ordination',
             'Bedöm anfallets längd på motorisk kramp och EEG-tid',
             'Stäm av läkemedlen inför serien'], size=20, h=4.3)

# Filmer: miniatyrbild som länk till TED.com och YouTube (TED-Ed-filmer har inte Creative Commons-licens och visas från källan)
def film_thumb(slide, x, y, w, image, url, descr):
    h = w * 9 / 16
    pic = slide.shapes.add_picture(str(IMG / image), Inches(x), Inches(y), Inches(w), Inches(h))
    pic.name = 'Film: ' + descr
    pic._element.nvPicPr.cNvPr.set('descr', 'Startbild för filmen ' + descr + '; länk till filmen')
    pic.click_action.hyperlink.address = url
    d = 0.9
    play = circle(slide, x + w - d - 0.2, y + h - d - 0.2, d, BLUE, [])
    play.name = 'Spela upp'
    play.click_action.hyperlink.address = url
    tri = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(x + w - d - 0.2 + 0.3), Inches(y + h - d - 0.2 + 0.25),
                                 Inches(0.4), Inches(0.38))
    tri.rotation = 90
    fill(tri, WHITE)
    tri.name = 'Spela upp, symbol'
    tri.click_action.hyperlink.address = url
    return h


TED_URL = 'https://www.ted.com/talks/helen_m_farrell_the_truth_about_electroconvulsive_therapy_ect'
s = add(L_MAX, 'Film: The truth about ECT')
film_thumb(s, 0.39, 1.65, 8.4, 'film-ted-ed.jpg', TED_URL, 'The truth about electroconvulsive therapy (ECT), TED-Ed')
box(s, 8.99, 1.65, 3.95, 4.73, BROWN, [('The truth about electroconvulsive therapy (ECT)', {'bold': True, 'size': 16, 'nospace': True}),
                                       'Helen M. Farrell, TED-Ed', 'Animerad film på engelska, 4 min 24 s',
                                       'Klicka på bilden för att öppna filmen på TED.com',
                                       ('© TED Conferences, LLC. TED-Ed-filmer har YouTubes standardlicens, inte Creative Commons, '
                                        'och visas direkt från TED.com eller YouTube.', {'size': 12})],
    size=14, space=8, inset=0.22)
source(s, 'Källa: Farrell HM. The truth about electroconvulsive therapy (ECT). TED-Ed; ted.com/talks/helen_m_farrell_the_truth_about_electroconvulsive_therapy_ect')

# Fler filmer
s = add(L_MAX, 'Fler filmer om ECT')
for x, vid, image, title, who in [
        (0.39, 'eP83fh2FswU', 'film-uc-davis.jpg', 'Electroconvulsive Therapy (ECT) Explained: How It Works, Recovery and Who It\'s For', 'UC Davis Health'),
        (6.74, '9L2-B-aluCE', 'film-st-georges.jpg', 'ECT', 'St George\'s, University of London & South West London & St George\'s Mental Health NHS Trust')]:
    url = f'https://www.youtube.com/watch?v={vid}'
    h = film_thumb(s, x, 1.55, 6.2, image, url, title)
    text(s, x, 1.55 + h + 0.12, 6.2, 1.6, [(title, {'bold': True, 'nospace': True}), who,
                                           (f'youtube.com/watch?v={vid}', {'size': 12})], size=14, space=4)
source(s, 'Filmerna öppnas på YouTube och tillhör respektive upphovsperson')

# 21 Frågor
s = add(L_SECTION, 'Frågor?')
sub = ph(s, 1)
sub._element.getparent().remove(sub._element)

prs.core_properties.title = 'ECT vid bordet'
prs.core_properties.subject = 'Stimulus, impedans, anfall och säkerhet'
notes = json.load(open(NOTES_FILE, encoding='utf-8'))
assert len(notes) == len(prs.slides), f'anteckningar.json har {len(notes)} poster, presentationen {len(prs.slides)} bilder'
for slide, note in zip(prs.slides, notes):
    if note:
        slide.notes_slide.notes_text_frame.text = note
prs.save(OUT)
embedded = embed(OUT)   # Noto Sans i stället för Verdana, inbäddat för PowerPoint (se typsnitt.py)
clean(OUT)   # inga personuppgifter eller organisationsdata från mallen (se metadata.py)
print('sparad', OUT, len(prs.slides), 'bilder; inbäddat typsnitt', ', '.join(f'{k} {v // 1024} kB' for k, v in embedded.items()))
